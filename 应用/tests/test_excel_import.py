"""End-to-end Excel import with disposable workbooks and isolated SQL data."""
import base64,io,zipfile
from datetime import datetime
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from openpyxl.styles import PatternFill
from test_business_regressions import client,db
import app as module
from excel_import import HEADERS

@pytest.fixture(autouse=True)
def cache_dir(monkeypatch,tmp_path):monkeypatch.setenv('IMPORT_DIR',str(tmp_path/'batches'))

def workbook(rows,extra=False):
 wb=Workbook();ws=wb.active;ws.title='我的原表';ws.append(HEADERS+(['备注'] if extra else []))
 for row in rows:ws.append(row)
 ws['A2'].fill=PatternFill('solid',fgColor='FFFF00')
 if extra:ws['H2']='保留原始备注';other=wb.create_sheet('原始说明');other['A1']='不要修改此表'
 out=io.BytesIO();wb.save(out);return out.getvalue()

def preview(client,raw,name='数据.xlsx'):
 return client.post('/dashboard/api/maintain/import/excel/preview',json={'file_name':name,'file_base64':base64.b64encode(raw).decode()})

def row(code='M71',day='2026-10-30',value=10,project='测试项目A',feature='',unit='人'):
 return [project,code,'原始指标名称',feature,day,value,unit]

def test_template_all_authorized_projects_and_dates(client,db):
 db.execute('INSERT INTO sys_role_project(role_id,project_id) VALUES(2,2)')
 response=client.get('/dashboard/api/maintain/import/template?stat_date=2026-10-30')
 assert response.status_code==200
 wb=load_workbook(io.BytesIO(response.content));ws=wb['数据导入']
 assert [c.value for c in ws[1]]==HEADERS
 assert {'测试项目A','测试项目B'}=={ws.cell(n,1).value for n in range(2,ws.max_row+1)}
 assert ws['E2'].value==datetime(2026,10,30)
 assert ws['F2'].value is None
 assert '填写说明' in wb.sheetnames
 assert client.get('/dashboard/api/maintain/import/template?stat_date=2099-01-01').status_code==400

def test_multi_project_multi_date_preview_upsert_and_original_report(client,db):
 db.execute('INSERT INTO sys_role_project(role_id,project_id) VALUES(2,2)')
 raw=workbook([row(),row(day='2026-10-29',value=20),row('M99',value=30,project='测试项目B'),row('M61',day='2026-10-29',value=11,feature='测试功能'),row('missing'),row(value=-3,day='2026-10-28'),row(value=None,day='2026-10-27')],True)
 response=preview(client,raw);assert response.status_code==200,response.text
 data=response.json();counts=data['counts']
 assert (counts['total'],counts['valid'],counts['invalid'],counts['insert'],counts['update'],counts['skipped'])==(7,4,2,3,1,1)
 # Preview validates without writing.
 assert db.execute('SELECT COUNT(*) FROM metric_entry').fetchone()[0]==2
 batch=data['batch_id'];url='/dashboard/api/maintain/import/excel/'+batch
 report=client.get(url+'/report');wb=load_workbook(io.BytesIO(report.content));ws=wb['我的原表']
 assert ws['H2'].value=='保留原始备注' and ws['A2'].fill.fgColor.rgb.endswith('FFFF00')
 assert wb['原始说明']['A1'].value=='不要修改此表'
 assert ws['I2'].value=='有效' and ws['J2'].value=='新增'
 assert ws['I6'].value=='无效' and '指标不存在' in ws['L6'].value
 result=client.post(url+'/commit');assert result.status_code==200,result.text
 data=result.json();assert data['phase']=='completed'
 assert (data['counts']['inserted'],data['counts']['updated'])==(3,1)
 assert db.execute('SELECT value_number FROM metric_entry WHERE id=1').fetchone()[0]==11
 assert client.post(url+'/commit').json()['counts']==data['counts']
 wb=load_workbook(io.BytesIO(client.get(url+'/report').content));assert wb['我的原表']['K2'].value=='已新增'
 history=client.get('/dashboard/api/maintain/history?project_id=1&q=Excel').json()['items']
 assert len(history)==3 and all(h['source']=='Excel 导入' for h in history)

def test_duplicate_rows_all_invalid_and_permissions_checked_per_project(client):
 data=preview(client,workbook([row(),row(),row('M99',project='测试项目B')])).json()
 assert data['counts']['invalid']==3 and data['counts']['valid']==0
 assert all('重复' in r['reason'] for r in data['rows'][:2])
 assert '没有导入权限' in data['rows'][2]['reason']

def test_cross_metric_validation_is_per_project_date_transaction(client,db):
 data=preview(client,workbook([row('M11',value=10),row('M12',value=20),row(day='2026-10-29',value=5)])).json()
 assert data['counts']['invalid']==2 and data['counts']['valid']==1
 assert all('累计使用用户数' in r['reason'] for r in data['rows'][:2])
 result=client.post('/dashboard/api/maintain/import/excel/'+data['batch_id']+'/commit').json()
 assert result['counts']['inserted']==1

def test_concurrent_update_rejected_and_other_date_imports(client,db):
 data=preview(client,workbook([row('M61',day='2026-10-29',value=11,feature='测试功能'),row(value=5)])).json()
 db.execute('UPDATE metric_entry SET value_number=12,row_version=2 WHERE id=1')
 result=client.post('/dashboard/api/maintain/import/excel/'+data['batch_id']+'/commit').json()
 assert result['counts']['invalid']==1 and result['counts']['inserted']==1 and result['counts']['updated']==0
 assert db.execute('SELECT value_number FROM metric_entry WHERE id=1').fetchone()[0]==12

def test_retention_observation_date_percentage_and_excel_dates(client,db):
 wb=Workbook();ws=wb.active;ws.append(HEADERS)
 ws.append(row('M31',value=.685,day=datetime(2026,10,30),unit='%'));ws['F2'].number_format='0.0%'
 out=io.BytesIO();wb.save(out)
 data=preview(client,out.getvalue()).json();assert data['counts']['valid']==1,data
 client.post('/dashboard/api/maintain/import/excel/'+data['batch_id']+'/commit')
 assert db.execute("SELECT value_number FROM metric_entry WHERE metric_id=31 AND stat_date='2026-10-23'").fetchone()[0]==68.5

def test_unknown_units_formula_future_precision_and_blank_zero(client):
 raw=workbook([row(value='=2+3'),row(day='2099-01-01'),row(day='2026-10-28',unit='次'),row(day='2026-10-27',value=.1),row(day='2026-10-26',value=0),row(day='2026-10-25',value=None),row(day='2026-10-24',value='1e999999999')])
 data=preview(client,raw).json();assert data['counts']['invalid']==5 and data['counts']['valid']==1 and data['counts']['skipped']==1,data
 assert data['rows'][4]['value']=='0'

def test_batch_owner_and_permission_revocation(client,db,monkeypatch):
 data=preview(client,workbook([row()])).json();url='/dashboard/api/maintain/import/excel/'+data['batch_id']
 monkeypatch.setattr(module,'actor',lambda r:{'id':1,'is_super':True})
 assert client.get(url+'/report').status_code==404
 monkeypatch.setattr(module,'actor',lambda r:{'id':2,'is_super':False})
 db.execute('DELETE FROM sys_role_project WHERE role_id=2 AND project_id=1')
 assert client.post(url+'/commit').status_code==403
 assert client.get(url+'/report').status_code==403

def test_invalid_file_and_limits(client):
 assert preview(client,b'not an xlsx').status_code==400
 assert preview(client,workbook([row()]),'macro.xlsm').status_code==400
 assert preview(client,workbook([row(day='2026-10-30')]*1001)).status_code==400
 assert client.post('/dashboard/api/maintain/import/excel/preview',json={'file_name':'a.xlsx','file_base64':'!'}).status_code==400

def test_no_import_permission_even_if_maintenance_editor(client,db):
 db.execute('DELETE FROM sys_role_permission WHERE permission_id=?',(db.permissions['maintain.import'],))
 assert client.get('/dashboard/api/maintain/import/template').status_code==403
 assert preview(client,workbook([row()])).status_code==403

def test_metric_configuration_changed_after_preview_rejected(client,db):
 data=preview(client,workbook([row()])).json()
 db.execute('UPDATE metric SET row_version=row_version+1,unit=? WHERE id=71',('秒',))
 result=client.post('/dashboard/api/maintain/import/excel/'+data['batch_id']+'/commit').json()
 assert result['counts']['inserted']==0 and result['counts']['invalid']==1
 assert '指标配置已变化' in result['rows'][0]['reason']

def test_expired_batch_and_delete_recreate_conflict(client,db,monkeypatch):
 import json,os,time
 data=preview(client,workbook([row('M61',day='2026-10-29',value=11,feature='测试功能')])).json()
 db.execute('DELETE FROM metric_entry WHERE id=1')
 db.execute("INSERT INTO metric_entry(project_id,metric_id,feature_id,stat_date,value_number) VALUES(1,61,1,'2026-10-29',12)")
 url='/dashboard/api/maintain/import/excel/'+data['batch_id']
 result=client.post(url+'/commit').json();assert result['counts']['updated']==0 and result['counts']['invalid']==1
 file=Path(os.environ['IMPORT_DIR'])/data['batch_id']/'result.json';batch=json.loads(file.read_text());batch['expires']=time.time()-1;file.write_text(json.dumps(batch))
 assert client.get(url+'/report').status_code==410

def test_duplicate_across_worksheets_and_unchanged_value(client,db):
 wb=Workbook();ws=wb.active;ws.append(HEADERS);ws.append(row())
 ws=wb.create_sheet('另一批');ws.append(HEADERS);ws.append(row())
 output=io.BytesIO();wb.save(output)
 assert preview(client,output.getvalue()).json()['counts']['invalid']==2
 data=preview(client,workbook([row('M61',day='2026-10-29',value=10,feature='测试功能')])).json()
 assert data['counts']['unchanged']==1
 result=client.post('/dashboard/api/maintain/import/excel/'+data['batch_id']+'/commit').json()
 assert result['counts']['inserted']==result['counts']['updated']==0
 assert db.execute('SELECT row_version FROM metric_entry WHERE id=1').fetchone()[0]==1
