"""Excel import workflow: authorized multi-project upsert and annotated original workbook.

The production Python service uses openpyxl for user-uploaded workbooks; no workbook
formulas are evaluated. Temporary batches are private to their uploader and expire.
"""
import base64,binascii,fcntl,io,json,os,re,secrets,shutil,time,zipfile
from collections import Counter,defaultdict
from contextlib import contextmanager
from datetime import date,datetime
from decimal import Decimal,InvalidOperation
from pathlib import Path
from urllib.parse import quote

from fastapi import HTTPException,Request
from fastapi.responses import Response
from pydantic import BaseModel,Field
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

HEADERS=['项目','指标编码','指标名称','功能','统计日期','数值','单位']
REQUIRED={'项目','指标编码','功能','统计日期','数值'}
ANNOTATIONS=['校验结果','处理类型','处理状态','原因说明','原值','导入值']
MAX_FILE=2*1024*1024
MAX_ROWS=1000
TTL=86400
MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

class ExcelFile(BaseModel):
    file_name:str=Field(min_length=1,max_length=200)
    file_base64:str=Field(min_length=1,max_length=2800000)

def excel_response(content,name):
    return Response(content,media_type=MIME,headers={'Content-Disposition':"attachment; filename*=UTF-8''"+quote(name)})

def save_workbook(wb):
    out=io.BytesIO();wb.save(out);return out.getvalue()

def text(cell,value):
    # Write labels/errors as literal text, never as generated spreadsheet formulas.
    cell.value=str(value) if value is not None else ''
    cell.data_type='s'

def style_table(ws,widths):
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(name='等线',bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor='245CE8');cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=30
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width

def open_excel(raw):
    if len(raw)>MAX_FILE:raise HTTPException(400,'Excel 文件不能超过 2 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            files=z.infolist()
            if len(files)>1000 or sum(f.file_size for f in files)>16*1024*1024:
                raise HTTPException(400,'Excel 解压内容过大，请精简文件后重试')
            if any('vbaproject' in f.filename.lower() for f in files):raise HTTPException(400,'请上传不含宏的 .xlsx 文件')
        wb=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
        if len(wb.worksheets)>10:raise HTTPException(400,'每个文件最多 10 个工作表')
        return wb
    except HTTPException:raise
    except Exception:raise HTTPException(400,'无法读取 Excel，请上传有效且未加密的 .xlsx 文件')

def number_value(cell,metric,S):
    if cell.data_type=='f':raise HTTPException(400,'数值不能使用公式，请粘贴为数值后导入')
    value=cell.value
    if isinstance(value,bool) or not isinstance(value,(int,float,str,Decimal)):
        raise HTTPException(400,'数值格式不正确')
    try:
        value=Decimal(str(value).strip())
        if not value.is_finite() or (value and value.adjusted()>17):raise InvalidOperation
        if value.as_tuple().exponent < -100:raise InvalidOperation
        if metric['unit']=='%' and '%' in cell.number_format:value*=100
        # Excel numeric 10.0 is an integer; trailing zeros do not change precision.
        raw=format(value,'f').rstrip('0').rstrip('.') if '.' in format(value,'f') else format(value,'f')
        return S._validate_value(raw or '0',metric)
    except (InvalidOperation,ValueError):raise HTTPException(400,'请输入非负数值')

def import_date(cell):
    if cell.data_type=='f':raise HTTPException(400,'统计日期不能使用公式')
    value=cell.value
    if isinstance(value,datetime):
        if value.time()!=datetime.min.time():raise HTTPException(400,'统计日期不应包含具体时间')
        return value.date()
    if isinstance(value,date):return value
    try:
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',str(value).strip()):raise ValueError
        return date.fromisoformat(str(value).strip())
    except ValueError:raise HTTPException(400,'统计日期应为 YYYY-MM-DD 或 Excel 日期')

def counts(rows):
    valid=[r for r in rows if r['validity']=='有效']
    actions=Counter(r['action'] for r in valid);states=Counter(r['status'] for r in rows)
    return dict(total=len(rows),valid=len(valid),invalid=sum(r['validity']=='无效' for r in rows),
        skipped=states['未填写'],insert=actions['新增'],update=actions['修改'],unchanged=actions['无需修改'],
        inserted=states['已新增'],updated=states['已修改'])

def public_batch(batch):
    return {'batch_id':batch['id'],'phase':batch['phase'],'file_name':batch['file_name'],
        'counts':counts(batch['rows']),'rows':[{k:r.get(k) for k in ['sheet','row','project','metric','date','value','old_value','validity','action','status','reason']} for r in batch['rows']]}

def register(S):
    root=lambda:Path(os.getenv('IMPORT_DIR',str(S.ROOT.parent/'imports')))
    def actor(request):
        u=S.actor(request);S.check_permission(u,'maintain.import');return u
    def scope(c,u):return S.scoped_projects(c,u,'maintain.import',True)
    def write_batch(path,batch):
        tmp=path/'result.tmp';tmp.write_text(json.dumps(batch,ensure_ascii=False));os.chmod(tmp,0o600);tmp.replace(path/'result.json')
    def cleanup():
        directory=root();directory.mkdir(parents=True,exist_ok=True,mode=0o700);os.chmod(directory,0o700)
        for path in directory.iterdir():
            if path.is_dir() and re.fullmatch(r'[A-Za-z0-9_-]{32}',path.name) and time.time()-path.stat().st_mtime>TTL+3600:shutil.rmtree(path)
        if sum(p.is_dir() for p in directory.iterdir())>=100:raise HTTPException(429,'临时导入文件较多，请稍后重试')
    @contextmanager
    def batch_access(batch_id,u):
        if not re.fullmatch(r'[A-Za-z0-9_-]{32}',batch_id):raise HTTPException(404,'导入记录不存在')
        path=root()/batch_id
        try:lock=(path/'lock').open('a')
        except FileNotFoundError:raise HTTPException(404,'导入记录不存在或已过期，请重新上传')
        with lock:
            fcntl.flock(lock,fcntl.LOCK_EX)
            try:batch=json.loads((path/'result.json').read_text())
            except (FileNotFoundError,ValueError):raise HTTPException(404,'导入记录不存在')
            if batch['owner']!=u['id']:raise HTTPException(404,'导入记录不存在')
            if time.time()>batch['expires']:raise HTTPException(410,'导入文件已过期，请重新上传')
            with S.database() as c:allowed={p['id'] for p in scope(c,u)}
            if not set(batch['project_ids'])<=allowed:raise HTTPException(403,'项目导入权限已变化，请重新上传校验')
            yield path,batch
    def payload(rows):
        return S.MaintainSave(project_id=rows[0]['project_id'],stat_date=rows[0]['date'],items=[
            S.MaintainItem(metric_id=r['metric_id'],feature_id=r['feature_id'],value=r['value'],expected_version=r['version'],expected_entry_id=r['entry_id'],expected_metric_version=r['metric_version']) for r in rows])
    def groups(rows):
        result=defaultdict(list)
        for r in rows:
            if r['validity']=='有效':result[(r['project_id'],r['date'])].append(r)
        return result

    @S.app.get(S.BASE+'/api/maintain/import/template')
    def template(request:Request,stat_date:str|None=None):
        u=actor(request)
        try:day=date.fromisoformat(stat_date) if stat_date else S._business_today()
        except ValueError:raise HTTPException(400,'模板日期格式无效')
        if day>S._business_today():raise HTTPException(400,'模板不能使用未来日期')
        wb=Workbook();ws=wb.active;ws.title='数据导入';ws.append(HEADERS)
        with S.database() as c:
            projects=scope(c,u)
            c.execute('''SELECT m.id,m.project_id,m.code,m.name,m.unit,m.decimal_places,m.row_version,cat.code AS category_code FROM metric m
                JOIN metric_category cat ON cat.id=m.category_id AND cat.is_enabled=1
                WHERE m.is_enabled=1 AND m.archived_at IS NULL ORDER BY m.display_order,m.id''');metrics=c.fetchall()
            c.execute('SELECT project_id,name FROM project_feature WHERE archived_at IS NULL ORDER BY display_order,id');features=c.fetchall()
        if not projects:raise HTTPException(403,'暂无有权导入的启用项目')
        for p in projects:
            for m in metrics:
                if m['project_id']!=p['id']:continue
                fs=[f['name'] for f in features if f['project_id']==p['id']] if m['category_code']=='feature' else ['']
                for f in fs:
                    if ws.max_row>MAX_ROWS:raise HTTPException(400,'全部项目模板超过 1,000 行，请联系管理员分批整理')
                    ws.append([p['name'],m['code'],m['name'],f,day,None,m['unit'] or ''])
                    for col in [1,2,3,4,7]:text(ws.cell(ws.max_row,col),ws.cell(ws.max_row,col).value)
                    ws.cell(ws.max_row,5).number_format='yyyy-mm-dd';ws.cell(ws.max_row,6).number_format='0'+('.'+'0'*m['decimal_places'] if m['decimal_places'] else '')
        style_table(ws,[24,20,30,24,18,18,18])
        help_ws=wb.create_sheet('填写说明');help_ws.append(['填写项','说明'])
        for row in [
            ('批量范围','同一文件可包含多个授权项目、多个统计日期；复制数据行后修改日期即可。'),
            ('数据定位','项目＋指标编码＋功能＋统计日期相同为同一条数据，已有数据会更新。'),
            ('指标与功能','保留项目、指标编码及对应功能；指标名称和单位用于核对，不改变指标配置。'),
            ('数值','直接填写非负数；比例填写 0～100（如 68.5 表示 68.5%），也支持 Excel 百分比单元格。空白数值跳过，不清空已有数据。'),
            ('日期','使用 YYYY-MM-DD 或 Excel 日期；不接受未来日期。留存统计日期为已结束的观察日。'),
            ('校验与导入','先点击导入校验，查看有效、无效、新增、修改统计，再确认导入有效数据。重复数据行全部无效。'),
            ('结果下载','在原数据工作表右侧标注校验结果、处理类型、状态、说明、原值与导入值；原始数据保持不变。'),
            ('限制','仅支持 .xlsx，2 MB、最多 1,000 条；不接受公式作为导入字段。回执保留 24 小时。'),
            ('事务范围','每个项目＋日期为一组；组内关系校验或并发冲突失败时整组不写入，其他有效组可导入。')]:help_ws.append(row)
        style_table(help_ws,[22,105])
        for row in help_ws.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='center')
            help_ws.row_dimensions[row[0].row].height=42
        return excel_response(save_workbook(wb),'指标批量导入模板.xlsx')

    @S.app.post(S.BASE+'/api/maintain/import/excel/preview')
    def preview(data:ExcelFile,request:Request):
        u=actor(request);S.mutation(request)
        if not data.file_name.lower().endswith('.xlsx'):raise HTTPException(400,'请上传 .xlsx 格式的 Excel 文件')
        try:raw=base64.b64decode(data.file_base64,validate=True)
        except (ValueError,binascii.Error):raise HTTPException(400,'文件编码不正确')
        wb=open_excel(raw);rows=[];duplicates=defaultdict(list);matched=False
        with S.database() as c:
            projects={p['name']:p for p in scope(c,u)}
            c.execute('''SELECT m.id,m.project_id,m.code,m.name,m.unit,m.decimal_places,m.row_version,cat.code AS category_code FROM metric m
                JOIN metric_category cat ON cat.id=m.category_id AND cat.is_enabled=1 WHERE m.is_enabled=1 AND m.archived_at IS NULL''')
            metrics={(m['project_id'],m['code']):m for m in c.fetchall()}
            c.execute('SELECT id,project_id,name FROM project_feature WHERE archived_at IS NULL')
            features={(f['project_id'],f['name']):f['id'] for f in c.fetchall()}
            for ws in wb.worksheets:
                if ws.max_column>64 or ws.max_row>10001:raise HTTPException(400,'工作表范围过大，请仅保留必要数据')
                headers=[str(cell.value or '').strip() for cell in ws[1]]
                if not REQUIRED<=set(headers):continue
                matched=True
                if any(headers.count(h)>1 for h in HEADERS):raise HTTPException(400,'导入列名重复，请重新下载模板')
                col={h:headers.index(h)+1 for h in HEADERS if h in headers}
                for n in range(2,ws.max_row+1):
                    cells={h:ws.cell(n,i) for h,i in col.items()}
                    if not any(cells[h].value not in [None,''] for h in REQUIRED):continue
                    r=dict(sheet=ws.title,row=n,project=str(cells['项目'].value or ''),metric=str(cells.get('指标名称',cells['指标编码']).value or ''),date=str(cells['统计日期'].value or ''),value='',old_value='',validity='无效',action='不导入',status='无效',reason='')
                    rows.append(r)
                    if len(rows)>MAX_ROWS:raise HTTPException(400,'每次最多导入 1,000 条数据')
                    if cells['数值'].value is None or str(cells['数值'].value).strip()=='':
                        r.update(validity='跳过',status='未填写',reason='数值为空，保留已有数据');continue
                    try:
                        if any(cell.data_type=='f' for cell in cells.values()):raise HTTPException(400,'导入字段不能使用公式，请粘贴为值')
                        p=projects.get(r['project'].strip())
                        if not p:raise HTTPException(400,'项目不存在、已停用或没有导入权限')
                        m=metrics.get((p['id'],str(cells['指标编码'].value or '').strip()))
                        if not m:raise HTTPException(400,'指标不存在、未启用或不属于该项目')
                        if '单位' in cells and str(cells['单位'].value or '').strip()!=(m['unit'] or ''):raise HTTPException(400,'单位与指标配置不一致，请重新下载模板')
                        fname=str(cells['功能'].value or '').strip();fid=features.get((p['id'],fname)) if fname else None
                        if m['category_code']=='feature' and fid is None:raise HTTPException(400,'请选择该项目的有效功能名称')
                        if m['category_code']!='feature' and fname:raise HTTPException(400,'该指标不支持按功能录入')
                        day=import_date(cells['统计日期'])
                        if day>S._business_today():raise HTTPException(400,'不能导入未来统计日')
                        if m['category_code']=='retention' and day>=S._business_today():raise HTTPException(400,'留存观察日必须已经结束')
                        value=number_value(cells['数值'],m,S)
                        stored=day-S.timedelta(days=7) if m['category_code']=='retention' else day
                        c.execute('SELECT id,value_number,row_version FROM metric_entry WHERE project_id=%s AND metric_id=%s AND feature_id<=>%s AND stat_date=%s',(p['id'],m['id'],fid,stored));old=c.fetchone()
                        old_value=S._fmt_decimal(old['value_number']) if old and old['value_number'] is not None else ''
                        action='新增' if not old else ('无需修改' if old_value!='' and Decimal(old_value)==Decimal(value) else '修改')
                        r.update(project_id=p['id'],metric_id=m['id'],feature_id=fid,date=day.isoformat(),metric=m['name'],value=value,old_value=old_value,version=old['row_version'] if old else 0,entry_id=old['id'] if old else None,metric_version=m['row_version'],validity='有效',action=action,status='待导入',reason='')
                        duplicates[(p['id'],m['id'],fid,day)].append(r)
                    except HTTPException as e:r['reason']=e.detail
        wb.close()
        if not matched:raise HTTPException(400,'未找到导入数据列，请使用下载的 Excel 模板')
        if not rows:raise HTTPException(400,'Excel 没有数据行')
        for group in duplicates.values():
            if len(group)>1:
                for r in group:r.update(validity='无效',action='不导入',status='无效',reason='文件内项目、指标、功能和日期重复，请仅保留一行')
        for group in groups(rows).values():
            try:S.save_maintenance(payload(group),request,'maintain.import','EXCEL',dry_run=True)
            except HTTPException as e:
                for r in group:r.update(validity='无效',status='无效',reason='同项目同日期校验未通过：'+str(e.detail))
        cleanup();batch_id=secrets.token_urlsafe(24);path=root()/batch_id;path.mkdir(mode=0o700)
        original=path/'original.xlsx';original.write_bytes(raw);os.chmod(original,0o600)
        batch=dict(id=batch_id,owner=u['id'],expires=time.time()+TTL,file_name=data.file_name,phase='preview',rows=rows,project_ids=sorted({r['project_id'] for r in rows if r['validity']=='有效'}))
        write_batch(path,batch);return public_batch(batch)

    @S.app.post(S.BASE+'/api/maintain/import/excel/{batch_id}/commit')
    def commit(batch_id:str,request:Request):
        u=actor(request);S.mutation(request)
        with batch_access(batch_id,u) as (path,batch):
            if batch['phase']=='completed':return public_batch(batch)
            if batch['phase']!='preview':raise HTTPException(409,'该批次正在处理或结果待核对，请下载回执后核对数据，不要重复提交')
            batch['phase']='processing';write_batch(path,batch)
            for group in groups(batch['rows']).values():
                try:
                    S.save_maintenance(payload(group),request,'maintain.import','EXCEL')
                    for r in group:r['status']={'新增':'已新增','修改':'已修改','无需修改':'无需修改'}[r['action']]
                except HTTPException as e:
                    for r in group:r.update(validity='无效',status='未导入',reason='导入失败：'+str(e.detail))
                except Exception:
                    for r in group:r.update(validity='无效',status='待核对',reason='服务异常，当前组未确认完成，请核对数据后重新校验')
                    batch['phase']='interrupted';write_batch(path,batch)
                    return public_batch(batch)
                write_batch(path,batch)
            batch['phase']='completed';write_batch(path,batch);return public_batch(batch)

    @S.app.get(S.BASE+'/api/maintain/import/excel/{batch_id}/report')
    def report(batch_id:str,request:Request):
        u=actor(request)
        with batch_access(batch_id,u) as (path,batch):
            wb=open_excel((path/'original.xlsx').read_bytes());starts={}
            for r in batch['rows']:
                ws=wb[r['sheet']]
                if ws.title not in starts:
                    start=ws.max_column+1;starts[ws.title]=start
                    for i,h in enumerate(ANNOTATIONS,start):
                        cell=ws.cell(1,i);text(cell,h);cell.fill=PatternFill('solid',fgColor='245CE8');cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True)
                        ws.column_dimensions[get_column_letter(i)].width=55 if h=='原因说明' else 18
                    ws.freeze_panes=ws.freeze_panes or 'A2'
                    ws.auto_filter.ref='A1:'+get_column_letter(start+len(ANNOTATIONS)-1)+str(ws.max_row)
                ws.row_dimensions[r['row']].height=max(ws.row_dimensions[r['row']].height or 15,45 if len(r['reason'])>28 else 28)
                color='FEE2E2' if r['validity']=='无效' else ('E0F2FE' if r['validity']=='有效' else 'F1F5F9')
                status=r['status']
                if batch['phase'] in ['processing','interrupted'] and status=='待导入':status='待核对'
                values=[r['validity'],r['action'],status,r['reason'],r['old_value'],r['value']]
                for i,v in enumerate(values,starts[ws.title]):
                    cell=ws.cell(r['row'],i);text(cell,v);cell.fill=PatternFill('solid',fgColor=color);cell.alignment=Alignment(vertical='top',wrap_text=True)
            title='本次导入统计';suffix=1
            while title in wb.sheetnames:suffix+=1;title='本次导入统计'+str(suffix)
            summary=wb.create_sheet(title);summary.append(['统计项','数量或状态'])
            c=counts(batch['rows'])
            for name,key in [('总记录','total'),('有效','valid'),('无效','invalid'),('数值未填写','skipped'),('新增','insert'),('修改','update'),('无需修改','unchanged'),('已新增','inserted'),('已修改','updated')]:summary.append([name,c[key]])
            summary.append(['处理阶段',{'preview':'待确认导入','completed':'处理完成','processing':'处理中，待核对','interrupted':'处理异常，待核对'}[batch['phase']]])
            style_table(summary,[26,35]);content=save_workbook(wb);wb.close()
            return excel_response(content,Path(batch['file_name']).stem+'_导入结果.xlsx')
